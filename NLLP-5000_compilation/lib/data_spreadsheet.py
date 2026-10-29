from lib import config
from lib.io_utils import dump_json, load_json
import csv
import logging
from openpyxl import load_workbook, Workbook
import os
import random
import re
import string


def _apply_rules_xlsx(headword_orig: str, pos_full: str, pos_new: str, row: int) -> dict:
    """Apply set of predefined rules to automatically populate inflected word forms in spreadsheet to be annotated.

    :param headword_orig: Original headword from Routledge dictionary.
    :param pos_full: Original, full POS tag from Routledge dictionary.
    :param pos_new: New POS tag.
    :param row: Row number.
    :return: Dictionary containing cell IDs as keys and cell content as values.
    """
    d_cells_to_vals = {}

    if pos_new in ["adj", "interj"]:
        d_cells_to_vals[f"F{row}"] = headword_orig

    if pos_new == "adj":

        if headword_orig[-3:-1] in ["aa", "ee", "oo", "uu"]:
            d_cells_to_vals[f"G{row}"] = f"{headword_orig[:-2]}{headword_orig[-1]}e"
            d_cells_to_vals[f"H{row}"] = f"{headword_orig}s"
            d_cells_to_vals[f"I{row}"] = f"{headword_orig[:-2]}{headword_orig[-1]}er"
            d_cells_to_vals[f"J{row}"] = f"{headword_orig[:-2]}{headword_orig[-1]}ere"
            d_cells_to_vals[f"K{row}"] = f"{headword_orig[:-2]}{headword_orig[-1]}ers"
        else:
            d_cells_to_vals[f"G{row}"] = f"{headword_orig}e"
            d_cells_to_vals[f"H{row}"] = f"{headword_orig}s" if not headword_orig.endswith(("s", "sch")) else headword_orig
            d_cells_to_vals[f"I{row}"] = f"{headword_orig}er"
            d_cells_to_vals[f"J{row}"] = f"{headword_orig}ere"
            d_cells_to_vals[f"K{row}"] = f"{headword_orig}ers"

        if headword_orig.endswith("sch"):  # words ending in "sch" do not have a superlative form
            d_cells_to_vals[f"L{row}"] = "N/A"
            d_cells_to_vals[f"M{row}"] = "N/A"
        else:
            d_cells_to_vals[f"L{row}"] = f"{headword_orig}st" if not headword_orig.endswith("s") else f"{headword_orig}t"
            d_cells_to_vals[f"M{row}"] = f"{headword_orig}ste" if not headword_orig.endswith("s") else f"{headword_orig}te"

    elif pos_new == "adv":
        d_cells_to_vals[f"I{row}"] = headword_orig

    elif pos_new == "art":
        d_cells_to_vals[f"H{row}"] = headword_orig

    elif pos_new == "conj":
        d_cells_to_vals[f"H{row}"] = headword_orig

    elif pos_new == "noun":

        if pos_full == "noun, de(m)":
            d_cells_to_vals[f"F{row}"] = "x"
        elif pos_full == "noun, het":
            d_cells_to_vals[f"G{row}"] = "x"
        elif pos_full == "noun, de(f)":
            d_cells_to_vals[f"H{row}"] = "x"
        elif pos_full in ["noun, de(m)/het", "noun, het/de(m)", "noun, de(m)|het"]:
            d_cells_to_vals[f"I{row}"] = "x"
        elif pos_full == "noun, de":
            d_cells_to_vals[f"K{row}"] = "x"
        elif pos_full == "noun, de(f)/het":
            d_cells_to_vals[f"L{row}"] = "x"
        else:
            pass

        d_cells_to_vals[f"N{row}"] = headword_orig

        if headword_orig[-3:-1] in ["aa", "ee", "oo", "uu"]:
            d_cells_to_vals[f"O{row}"] = f"{headword_orig[:-2]}{headword_orig[-1]}en"
        else:
            d_cells_to_vals[f"O{row}"] = f"{headword_orig}en" if not headword_orig.endswith("e") else f"{headword_orig}n"

        if headword_orig.endswith("ing"):
            d_cells_to_vals[f"P{row}"] = f"{headword_orig[:-1]}kje"
            d_cells_to_vals[f"Q{row}"] = f"{headword_orig[:-1]}kjes"
        elif headword_orig.endswith(("d", "f", "g", "h", "k", "p", "s", "t")):
            d_cells_to_vals[f"P{row}"] = f"{headword_orig}je"
            d_cells_to_vals[f"Q{row}"] = f"{headword_orig}jes"
        else:
            d_cells_to_vals[f"P{row}"] = f"{headword_orig}tje"
            d_cells_to_vals[f"Q{row}"] = f"{headword_orig}tjes"

    elif pos_new == "num":
        d_cells_to_vals[f"J{row}"] = headword_orig

    elif pos_new == "prep":
        d_cells_to_vals[f"I{row}"] = headword_orig

    elif pos_new == "verb":
        d_cells_to_vals[f"J{row}"] = headword_orig
        d_cells_to_vals[f"N{row}"] = headword_orig

    else:
        pass

    return d_cells_to_vals


def _load_d_headers_spreadsheet() -> dict:
    """Get the headers to be used in the spreadsheet.

    :return: A dictionary containing the headers per POS.
    """
    with open(config.DIREC_SPREADSHEET_HEADERS / "common.txt", mode="r", encoding="utf-8") as f:
        l_common_headers = [header.strip() for header in f.readlines()]

    l_pos = [
        doc.replace(".txt", "").replace("specific_", "")
        for doc in os.listdir(config.DIREC_SPREADSHEET_HEADERS) if doc.startswith("specific_")
    ]
    d_headers = {}

    for pos in l_pos:

        with open(config.DIREC_SPREADSHEET_HEADERS / f"specific_{pos}.txt", mode="r", encoding="utf-8") as f:
            l_headers_pos = [header.strip() for header in f.readlines()]

        d_headers[pos] = l_common_headers + l_headers_pos

    return d_headers


def _load_d_mapping_pos_orig_to_new(d_headwords_orig: dict) -> dict:
    """Map part-of-speech (POS) tags as used in the Routledge resource to new POS tag (e.g., mapping masculine nouns,
    feminine nouns, and neuter nouns to the overarching tag 'noun').

    :param d_headwords_orig: Structured dictionary of the raw frequency dictionary.
    :return: The mapping dictionary.
    """
    d_mapping = {}

    for word in d_headwords_orig:

        for pos in d_headwords_orig[word]["POS_tags"]:

            if pos not in d_mapping:

                if pos in {"art", "conj", "prep", "verb", "pron", "adv", "num", "adj", "interj"}:
                    d_mapping[pos] = pos
                elif "noun" in pos:
                    d_mapping[pos] = "noun"
                else:
                    raise ValueError(f"Invalid POS tag: {pos} (headword = {word}).")

    return d_mapping


def convert_raw_input_into_dic(path_input_data: str, logger: logging.Logger) -> dict:
    """Convert raw input data into a structured Python dictionary.

    :param path_input_data: Path where the file containing the original Routledge dictionary in TXT format is located.
    :param logger: The logger.
    :return: The structured dictionary.
    """
    # LOAD RAW FREQUENCY LIST
    with open(path_input_data, mode="r", encoding="utf-8") as f:
        l_lines = [tuple(line.strip().split("\t")) for line in f.readlines()]

    # LOOP OVER LINES FROM RAW FILE AND STORE INFORMATION IN STRUCTURED DICTIONARY
    d_headwords_orig = {}
    l_headwords_orig = []  # list of headwords to be used for calculating number of duplicate entries (i.e. entries that occur in more than one list)
    d_sense_id_to_letter = {idx: letter for idx, letter in enumerate(string.ascii_lowercase, start=1)}

    for idx_line, tup_line in enumerate(l_lines):

        # Check if line is special line (i.e. line that does not contain an entry of the frequency list)
        if len(tup_line) != 11:
            logger.info(f"Line {idx_line + 1} does not contain 11 elements (only {len(tup_line)} found) and "
                        f"will be skipped: {tup_line}.")
            continue

        # Retrieve relevant information from line
        headword_orig = tup_line[1]  # original Routledge headword is situated at position 1 in tuple
        pos_raw = tup_line[2]  # Routledge POS is situated at position 2 (raw, unprocessed version)
        translation_en_raw = tup_line[3]  # English translation of the headword is situated at position 3
        tup_example_sents = tup_line[4:8]  # example sentences are situated at positions 4 to 7
        tup_mwes = tup_line[8:10]  # MWEs are situated at positions 8 to 9
        freq = float(tup_line[-1])  # frequency value is situated at the last position
        l_headwords_orig.append((headword_orig, pos_raw, freq))

        # Create empty dictionary in which data will be stored per POS
        d_data_per_pos = {}

        # Process entries that have more than one POS assigned to them (indicated in list by means of "1) ", "2) ",
        # etc.) for these entries, a separate entry per POS is created in the dictionary
        if "1) " in pos_raw:
            logger.info(f"\t\tEntry with more than one POS assigned to it: {tup_line}.")
            l_ids_pos = [f"POS_{item.strip(') ')}" for item in re.findall("[0-9]\\) ", pos_raw)]
            l_pos = [item.strip() for item in re.split("[0-9]\\) ", pos_raw) if item != ""]
            l_translations_en = [item.strip() for item in re.split("[0-9]\\) ", translation_en_raw) if item != ""]
            assert len(l_ids_pos) == len(l_pos) == len(l_translations_en)

            for id_pos, pos in zip(l_ids_pos, l_pos):

                if id_pos not in d_data_per_pos:
                    d_data_per_pos[id_pos] = {"POS": pos, "senses": {"sense_1": {}}}  # entries with multiple POSs never included multiple senses in the original Routledge dictionary (manually verified), so an empty sense entry can already be created

            for id_pos, translation_en in zip(l_ids_pos, l_translations_en):
                assert "a) " not in translation_en  # double check that entry does not have more than one sense
                d_data_per_pos[id_pos]["senses"]["sense_1"]["translation_EN"] = translation_en

        # Process entries that have one POS assigned to them
        else:

            #   - Retrieve different translations of the entry (indicated by means of "a) ", "b) ", etc.)
            l_ids_translations_en = [
                item.strip(") ") for item in re.findall("(^a\\) | [b-z]\\) )", translation_en_raw)
            ]
            l_translations_en = [
                item.strip() for item in re.split("^a\\) | [b-z]\\) ", translation_en_raw) if item != ""
            ]
            assert len(l_ids_translations_en) == 0 or len(l_ids_translations_en) >= 2, tup_line  # make sure no undesired matches are identified (either no matches are found or at least two)

            #   - Store translations to dictionary (as different senses)
            if len(l_ids_translations_en) == 0:  # only one translation = only one sense (length of `l_ids_translations_en` = 0 because no sense separators were found)
                d_data_per_pos["POS_1"] = {
                    "POS": pos_raw.strip(), "senses": {"sense_1": {"translation_EN": translation_en_raw.strip()}}
                }
            else:  # multiple translations = multiple senses (length of `l_ids_translations_en` = number of senses)
                d_data_per_pos["POS_1"] = {"POS": pos_raw.strip(), "senses": {}}

                for id_sense, translation_en in enumerate(l_translations_en):
                    entry_sense = f"sense_{id_sense + 1}"
                    assert entry_sense not in d_data_per_pos["POS_1"]["senses"]
                    d_data_per_pos["POS_1"]["senses"][entry_sense] = {"translation_EN": translation_en}

        if headword_orig not in d_headwords_orig:
            d_headwords_orig[headword_orig] = {"POS_tags": {}}

        for id_pos in d_data_per_pos:
            pos = d_data_per_pos[id_pos]["POS"]

            for id_sense in d_data_per_pos[id_pos]["senses"]:
                translation_en = d_data_per_pos[id_pos]["senses"][id_sense]["translation_EN"]

                if len(d_data_per_pos) > 1:
                    assert len(d_data_per_pos[id_pos]["senses"]) == 1  # double check that entries with multiple POSs have only one sense
                    l_example_sents = [
                        sent[3:] for sent in tup_example_sents if sent.startswith(f"{id_pos.split('_')[-1]}) ")
                    ]  # select example sentence corresponding to POS from tuple
                else:

                    if len(d_data_per_pos[id_pos]["senses"]) > 1:
                        l_example_sents = [
                            sent[3:] for sent in tup_example_sents
                            if sent.startswith(f"{d_sense_id_to_letter[int(id_sense.split('_')[-1])]}) ")
                        ]  # select example sentence corresponding to sense from tuple
                    else:
                        l_example_sents = [sent for sent in tup_example_sents if sent != ""]

                if pos not in d_headwords_orig[headword_orig]["POS_tags"]:
                    d_headwords_orig[headword_orig]["POS_tags"][pos] = {"senses": {}, "freq": freq}
                else:
                    assert d_headwords_orig[headword_orig]["POS_tags"][pos]["freq"] == freq, (headword_orig, pos, d_headwords_orig[headword_orig], freq)  # double check that words with multiple occurrences in word list have same frequency assigned to them

                if translation_en not in d_headwords_orig[headword_orig]["POS_tags"][pos]["senses"]:
                    d_headwords_orig[headword_orig]["POS_tags"][pos]["senses"][translation_en] = []

                for sent in l_example_sents:
                    d_headwords_orig[headword_orig]["POS_tags"][pos]["senses"][translation_en].append(sent)

    # DETERMINE NUMBER OF DUPLICATE ENTRIES
    n_duplicates = len(l_headwords_orig) - len(set(l_headwords_orig))
    logger.info(f"Total number of entries in original resource = {len(l_headwords_orig)} | "
                f"Number of unique entries = {len(set(l_headwords_orig))} | "
                f"Number of duplicate entries = {n_duplicates}")

    # DUMP DICTIONARY TO JSON FILE FOR FUTURE REFERENCE
    os.makedirs(config.DIREC_INP_DATA_PROCSD, exist_ok=True)
    dump_json(config.DIREC_INP_DATA_PROCSD, "d_headwords_orig.json", d_headwords_orig, indent=2)

    return d_headwords_orig


def _process_row(pos_new: str, row: tuple, logger: logging.Logger) -> tuple[int, str, str, str, list, list]:
    """Process a row from the annotated XLSX file.

    :param pos_new: New POS tag.
    :param row: The spreadsheet row containing the annotations.
    :param logger: The logger.
    :return: A tuple containing the frequency value, the headword, the original full POS tag, the checked annotation,
        the list of raw annotations, and the list of inflected/derived forms + UD POS + UD morphological features.
    """
    # EXTRACT FREQUENCY VALUE, HEADWORD AND ORIGINAL, FULL POS TAG FROM ROW
    freq = row[2].value if "|" not in str(row[2].value) else max([float(v) for v in row[2].value.split("|")])
    headword = row[3].value
    pos_full = row[4].value

    # EXTRACT RAW ANNOTATIONS FROM ROW AND DEFINE INFLECTED/DERIVED FORMS (+ EXTRA UD DATA) BASED ON ANNOTATIONS
    l_inflections = []

    if pos_new == "adj":

        # ===== Raw annotations =====
        l_raw_annots = [
            row[5].value, row[6].value, row[7].value, row[8].value, row[9].value, row[10].value, row[11].value,
            row[12].value
        ]

        # ===== Inflections and derivations =====
        d_mapping_cols_feats = {
            5: "Degree=Pos", 6: "Degree=Pos", 7: "Degree=Pos",
            8: "Degree=Cmp", 9: "Degree=Cmp", 10: "Degree=Cmp",
            11: "Degree=Sup", 12: "Degree=Sup"
        }  # dictionary mapping column numbers to UD morphological features of inflected/derived word forms stored in the respective columns

        for col in d_mapping_cols_feats:

            if row[col].value != "N/A":  # skip non-existing inflections/derivations
                l_inflections.append({
                    "word_form": row[col].value, "POS_UD": "ADJ", "morphological_features_UD": d_mapping_cols_feats[col]
                })

    elif pos_new == "adv":
        invariable_word_form = row[8].value

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            invariable_word_form
        ]

        # ===== Inflections and derivations =====
        if row[5].value == "x":  # regular adverbs
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADV", "morphological_features_UD": None}
            )
        elif row[6].value == "x" and row[7].value != "x":  # negative adverbs
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADV", "morphological_features_UD": "PronType=Neg"}
            )
        elif row[6].value == "x" and row[7].value == "x":  # particles encoding negation (e.g., "niet")
            l_inflections.append(
                {"word_form": invariable_word_form,
                 "POS_UD": "PART", "morphological_features_UD": "PartType=Neg|Polarity=Neg"}
            )
        elif row[6].value != "x" and row[7].value == "x":  # verbal particles
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADP", "morphological_features_UD": "PartType=Vbp"}
            )
        else:
            raise ValueError(f"Invalid annotation for {headword}: check X marks.")

    elif pos_new == "art":

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False, True if row[6].value == "x" else False,
            row[7].value, row[8].value, row[9].value, row[10].value, row[11].value
        ]

        # ===== Inflections and derivations =====
        feat_definite = None

        if row[5].value == "x":  # definite articles
            assert row[6].value != "x"
            feat_definite = "Definite=Def"

        if row[6].value == "x":  # indefinite articles
            assert row[5].value != "x"
            feat_definite = "Definite=Ind"

        assert feat_definite is not None, f"No definiteness marked for {headword}."

        d_mapping_cols_feats = {7: "Case=Nom", 8: "Case=Acc", 9: "Case=Gen", 10: "Case=Dat"}  # dictionary mapping column numbers to UD morphological features of inflected/derived word forms stored in the respective columns

        for col in d_mapping_cols_feats:

            if row[col].value != "N/A":  # skip non-existing inflections/derivations
                l_inflections.append({
                    "word_form": row[col].value, "POS_UD": "DET",
                    "morphological_features_UD": f"PronType=Art|{feat_definite}|{d_mapping_cols_feats[col]}"
                })

        # Additional column containing possible inflections with "buigings-e", needs special handling
        d_mapping_case_feats = {"nominatief": "Case=Nom", "accusatief": "Case=Acc", "datief": "Case=Dat"}

        if row[11].value != "N/A":  # skip non-existing inflections

            for item in row[11].value.split(";"):
                feat_case = d_mapping_case_feats[item.split("=")[0]]
                word_form = item.split("=")[1]
                l_inflections.append({
                    "word_form": word_form, "POS_UD": "DET",
                    "morphological_features_UD": f"PronType=Art|{feat_definite}|{feat_case}"
                })

    elif pos_new == "conj":
        invariable_word_form = row[7].value

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False, True if row[6].value == "x" else False,
            invariable_word_form
        ]

        # ===== Inflections and derivations =====
        if row[5].value == "x":  # coordinating conjunctions
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "CCONJ", "morphological_features_UD": None}
            )

        if row[6].value == "x":  # subordinating conjunctions
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "SCONJ", "morphological_features_UD": None}
            )

    elif pos_new == "interj":

        # ===== Raw annotations =====
        l_raw_annots = [row[5].value]

        # ===== Inflections and derivations =====
        l_inflections.append({"word_form": row[5].value, "POS_UD": "INTJ", "morphological_features_UD": None})

    elif pos_new == "noun":

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            True if row[8].value == "x" else False,
            True if row[9].value == "x" else False,
            True if row[10].value == "x" else False,
            True if row[11].value == "x" else False,
            True if row[12].value == "x" else False,
            row[13].value, row[14].value, row[15].value, row[16].value
        ]

        # ===== Inflections and derivations =====
        if row[12].value == "x":
            pos_ud = "PROPN"
        else:
            pos_ud = "NOUN"

        feat_gender = None

        if row[5].value == "x":
            assert row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Masc"

        if row[6].value == "x":
            assert row[5].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Neut"

        if row[7].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Fem"

        if row[8].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Masc,Neut"

        if row[9].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Fem,Masc,Neut"

        if row[10].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[11].value != "x"
            feat_gender = "Gender=Fem,Masc"

        if row[11].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x"
            feat_gender = "Gender=Fem,Neut"

        d_mapping_cols_feats = {
            13: "Number=Sing", 14: "Number=Plur", 15: "Number=Sing|Degree=Dim", 16: "Number=Plur|Degree=Dim"
        }

        for col in d_mapping_cols_feats:

            if row[col].value != "N/A":  # skip non-existing inflections/derivations
                morphological_features_ud = d_mapping_cols_feats[col]

                if feat_gender is not None:  # if statement needed because some proper nouns do not hava a gender
                    morphological_features_ud += f"|{feat_gender}"

                l_inflections.append({
                    "word_form": row[col].value, "POS_UD": pos_ud,
                    "morphological_features_UD": morphological_features_ud
                })

    elif pos_new == "num":

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            True if row[8].value == "x" else False,
            row[9].value, row[10].value, row[11].value
        ]

        # ===== Inflections and derivations =====
        feat_numtype = None
        feat_definite = None

        if row[5].value == "x":
            assert row[6].value != "x"
            feat_numtype = "NumType=Card"

        if row[6].value == "x":
            assert row[5].value != "x"
            feat_numtype = "NumType=Ord"

        if row[7].value == "x":
            assert row[8].value != "x"
            feat_definite = "Definite=Def"

        if row[8].value == "x":
            assert row[7].value != "x"
            feat_definite = "Definite=Ind"

        assert feat_numtype is not None and feat_definite is not None

        for word_form in [row[9].value, row[10].value, row[11].value]:

            if word_form != "N/A":  # skip non-existing inflections/derivations
                l_inflections.append({
                    "word_form": word_form, "POS_UD": "NUM",
                    "morphological_features_UD": f"{feat_numtype}|{feat_definite}"
                })

    elif pos_new == "prep":
        invariable_word_form = row[8].value

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            invariable_word_form
        ]

        # ===== Inflections and derivations =====
        if row[5].value == "x":  # postpositions
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADP", "morphological_features_UD": "AdpType=Post"}
            )

        if row[6].value == "x":  # circumpositions
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADP", "morphological_features_UD": "AdpType=Circ"}
            )

        if row[7].value == "x":  # prepositions
            l_inflections.append(
                {"word_form": invariable_word_form, "POS_UD": "ADP", "morphological_features_UD": "AdpType=Prep"}
            )

    elif pos_new == "pron":

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            True if row[8].value == "x" else False,
            True if row[9].value == "x" else False,
            True if row[10].value == "x" else False,
            True if row[11].value == "x" else False,
            True if row[12].value == "x" else False,
            row[13].value, row[14].value, row[15].value, row[16].value, row[17].value, row[18].value, row[19].value,
            row[20].value, row[21].value,
        ]

        # ===== Inflections and derivations =====
        feat_prontype = None

        if row[5].value == "x":
            assert row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Dem"

        if row[6].value == "x":
            assert row[5].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Rel"

        if row[7].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Prs|Poss=Yes"

        if row[8].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Ind"

        if row[9].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[10].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Prs"

        if row[10].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[11].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Prs|Reflex=Yes"

        if row[11].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[12].value != "x"
            feat_prontype = "PronType=Rcp"

        if row[12].value == "x":
            assert row[5].value != "x" and row[6].value != "x" and row[7].value != "x" and row[8].value != "x" and row[9].value != "x" and row[10].value != "x" and row[11].value != "x"
            feat_prontype = "PronType=Int"

        if row[13].value != "N/A":
            assert row[13].value in ["Sing", "Plur"]
            feat_number = f"Number={row[13].value}"
        else:
            feat_number = None

        if row[14].value != "N/A":
            assert row[14].value in [1, 2, 3]
            feat_person = f"Person={row[14].value}"
        else:
            feat_person = None

        if row[15].value != "N/A":
            assert row[15].value in ["Fem", "Masc", "Neut", "Fem,Masc", "Fem,Masc,Neut", "Fem,Neut", "Masc,Neut"]
            feat_gender = f"Gender={row[15].value}"
        else:
            feat_gender = None

        assert feat_prontype is not None, f"No pronoun type marked for {headword}."

        morph_feats_ud = feat_prontype

        for feat in [feat_number, feat_person, feat_gender]:

            if feat is not None:
                morph_feats_ud += f"|{feat}"

        d_mapping_cols_feats = {16: "Case=Nom", 17: "Case=Acc", 18: "Case=Gen", 19: "Case=Dat"}

        for col in d_mapping_cols_feats:

            if row[col].value != "N/A":  # skip non-existing inflections/derivations
                l_inflections.append({
                    "word_form": row[col].value, "POS_UD": "PRON",
                    "morphological_features_UD": f"{morph_feats_ud}|{d_mapping_cols_feats[col]}"
                })

        d_mapping_case_feats = {"nominatief": "Case=Nom", "accusatief": "Case=Acc", "datief": "Case=Dat"}

        # Additional columns containing possible inflections with "buigings-e" and "buigings-n", needs special handling
        for cell_val in [row[20].value, row[21].value]:

            if cell_val != "N/A":  # skip non-existing inflections

                for item in cell_val.split(";"):
                    feat_case = d_mapping_case_feats[item.split("=")[0]]
                    word_form = item.split("=")[1]
                    l_inflections.append({
                        "word_form": word_form, "POS_UD": "PRON",
                        "morphological_features_UD": f"{morph_feats_ud}|{feat_case}"
                    })

    elif pos_new == "verb":
        special_verb = True if row[0].value == "SPECIAL" else False

        # ===== Raw annotations =====
        l_raw_annots = [
            True if row[5].value == "x" else False,
            True if row[6].value == "x" else False,
            True if row[7].value == "x" else False,
            True if row[8].value == "x" else False,
            row[9].value, row[10].value, row[11].value, row[12].value, row[13].value, row[14].value, row[15].value,
            row[16].value, row[17].value, row[18].value, row[19].value,
        ]

        # ===== Inflections and derivations =====
        l_pos_feat_verbtype_ud = []

        if row[5].value == "x":
            l_pos_feat_verbtype_ud.append(("VERB", None))

        if row[6].value == "x":
            l_pos_feat_verbtype_ud.append(("AUX", None))

        if row[7].value == "x":
            l_pos_feat_verbtype_ud.append(("AUX", "VerbType=Cop"))

        if row[8].value == "x":
            l_pos_feat_verbtype_ud.append(("AUX", "VerbType=Mod"))

        d_mapping_cols_feats = {
            9: ["VerbForm=Inf"],
            10: [
                "VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Sing|Person=1",
                "VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Sing|Person=2"
            ],
            11: ["VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Sing|Person=2"],
            12: ["VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Sing|Person=3"],
            13: [
                "VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Plur|Person=1",
                "VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Plur|Person=2",
                "VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Plur|Person=3"
            ] if not special_verb else ["VerbForm=Fin|Mood=Ind|Tense=Pres|Number=Plur|Person=3"],
            14: [
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Sing|Person=1",
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Sing|Person=2",
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Sing|Person=3"
            ] if not special_verb else ["VerbForm=Fin|Mood=Ind|Tense=Past|Number=Sing|Person=3"],
            15: [
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Plur|Person=1",
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Plur|Person=2",
                "VerbForm=Fin|Mood=Ind|Tense=Past|Number=Plur|Person=3"
            ] if not special_verb else ["VerbForm=Fin|Mood=Ind|Tense=Past|Number=Plur|Person=3"],
            16: ["VerbForm=Fin|Mood=Imp"],
            17: ["VerbForm=Fin|Mood=Sub"],
            18: ["VerbForm=Part|Tense=Pres"],
            19: ["VerbForm=Part|Tense=Past"]
        }

        for tup in l_pos_feat_verbtype_ud:
            pos_ud = tup[0]
            feat_verbtype = tup[1]

            for col in d_mapping_cols_feats:

                for morph_feats_ud in d_mapping_cols_feats[col]:

                    if feat_verbtype is not None:
                        morph_feats_ud += f"|{feat_verbtype}"

                    if row[col].value != "N/A":  # skip non-existing inflections/derivations
                        l_inflections.append({
                            "word_form": row[col].value, "POS_UD": pos_ud, "morphological_features_UD": morph_feats_ud
                        })

    else:
        raise ValueError(f"Unknown POS: {pos_new} (headword = {headword}.")

    return freq, headword, pos_full, row[0].value, l_raw_annots, l_inflections


def convert_d_headwords_orig_into_unverified_spreadsheet(d_headwords_orig: dict) -> None:
    """Convert structured dictionary list into spreadsheet (CSV and XLSX) to be verified and completed by expert (some
    fields are automatically populated beforehand based on a set of predefined rules).

    :param d_headwords_orig: Structured dictionary.
    :return: `None`
    """
    # LOAD SUPPORT DATA
    d_headers = _load_d_headers_spreadsheet()
    d_mapping_pos = _load_d_mapping_pos_orig_to_new(d_headwords_orig)
    l_pos_new_sorted = sorted(list(set(list(d_mapping_pos.values()))))

    # CREATE EMPTY SPREADSHEET

    #   - XLSX
    os.makedirs(config.DIREC_INP_DATA_PROCSD, exist_ok=True)
    wb = Workbook()
    del wb["Sheet"]  # delete automatically created "Sheet" worksheet

    for pos_new in l_pos_new_sorted:
        ws = wb.create_sheet(pos_new)
        column = 1

        for header in d_headers[pos_new]:
            ws.cell(row=1, column=column, value=header)
            column += 1

    d_idx_row = {pos_new: 2 for pos_new in l_pos_new_sorted}

    for headword_orig in d_headwords_orig:

        for pos_word_full in d_headwords_orig[headword_orig]["POS_tags"]:
            pos_word_new = d_mapping_pos[pos_word_full]
            ws = wb[pos_word_new]
            row = d_idx_row[pos_word_new]
            ws[f"C{row}"] = d_headwords_orig[headword_orig]["POS_tags"][pos_word_full]["freq"]
            ws[f"D{row}"] = headword_orig
            ws[f"E{row}"] = pos_word_full

            if "," not in headword_orig:
                d_cells_to_vals = _apply_rules_xlsx(headword_orig, pos_word_full, pos_word_new, row)

                for cell in d_cells_to_vals:
                    ws[cell] = d_cells_to_vals[cell]

            d_idx_row[pos_word_new] += 1

    wb.save(os.path.join(config.DIREC_INP_DATA_PROCSD, "NLLP-5000_inflectedForms_unverified.xlsx"))

    #   - CSVs
    os.makedirs(config.DIREC_INP_DATA_PROCSD / "NLLP-5000_inflectedForms_unverified_CSVs", exist_ok=True)
    d_rows_per_pos = {pos_new: [] for pos_new in l_pos_new_sorted}
    l_col_letters_xlsx = [
        "A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V",
        "W", "X", "Y", "Z"
    ]

    for headword_orig in d_headwords_orig:

        for pos_word_full in d_headwords_orig[headword_orig]["POS_tags"]:
            pos_word_new = d_mapping_pos[pos_word_full]
            d_data_xlsx_to_csv = {letter: "" for letter in l_col_letters_xlsx[:len(d_headers[pos_word_new])]}
            d_data_xlsx_to_csv["C"] = d_headwords_orig[headword_orig]["POS_tags"][pos_word_full]["freq"]
            d_data_xlsx_to_csv["D"] = headword_orig
            d_data_xlsx_to_csv["E"] = pos_word_full

            if "," not in headword_orig:
                d_cells_to_vals = _apply_rules_xlsx(headword_orig, pos_word_full, pos_word_new, 0)  # row is irrelevant, so put to 0

                for cell in d_cells_to_vals:
                    d_data_xlsx_to_csv[cell[0]] = d_cells_to_vals[cell]  # `cell[0]` corresponds to column letter XLSX

            if pos_word_new not in d_rows_per_pos:
                d_rows_per_pos[pos_word_new] = []

            d_rows_per_pos[pos_word_new].append(tuple([d_data_xlsx_to_csv[letter] for letter in l_col_letters_xlsx[:len(d_headers[pos_word_new])]]))

    for pos in l_pos_new_sorted:

        with open(
                config.DIREC_INP_DATA_PROCSD / "NLLP-5000_inflectedForms_unverified_CSVs" / f"NLLP-5000_inflectedForms_unverified_{pos}.csv",
                mode="w", newline="", encoding="utf-8"
        ) as f_csv:
            csv_writer = csv.writer(f_csv, delimiter=",", quotechar='"', quoting=csv.QUOTE_MINIMAL)
            csv_writer.writerow(d_headers[pos])

            for tup in d_rows_per_pos[pos]:
                csv_writer.writerow(tup)


def convert_verified_spreadsheet_into_nllp5000(path_xlsx_filled_in: str, logger: logging.Logger) -> None:
    """Convert spreadsheet (XLSX format) containing the verified and completed inflected forms into the final NLLP-5000
    resource.

    :param path_xlsx_filled_in: Path where the XLSX file containing the verified and completed inflected forms is
        located.
    :param logger: The logger.
    :return: `None`
    """
    # PROCESS ANNOTATIONS
    d_headword_to_freq = {}  # dictionary in which headword (i.e. base form linked to a POS = lemma) is linked to frequency value, needed to determine frequency rank of the headwords
    l_entries_prov = []  # list of provisional entries (ID, frequency rank, and frequency band still need to be added later on)
    n_inflected_forms_full = 0  # NOTE: This variable also includes derived forms, the term "inflected" is chosen for ease of reference. The same applies to other variable names including "inflections/inflected".
    wb = load_workbook(filename=path_xlsx_filled_in)

    for pos in wb.sheetnames:
        logger.info(f"Running script for {pos} ...")
        l_tups = []
        l_entries_excluded = []

        for row in wb[pos].iter_rows(min_row=2):
            checked = row[0].value

            if checked in ["x", "CHANGED", "NEW", "MOVED", "SPECIAL"]:  # only consider rows that have been checked ("x") or that received a special annotation (other options)
                l_tups.append(_process_row(pos, row, logger))
            else:
                l_entries_excluded.append((pos, checked, row[2].value, row[3].value))

        logger.info(f"\t\tNumber of entries excluded for {pos} = {len(l_entries_excluded)}. "
                    f"First five entries (if applicable): {l_entries_excluded[:5]}.\n")

        for tup in l_tups:
            freq = tup[0]
            word = tup[1]
            l_inflections = tup[5]
            assert l_inflections, f"No inflected forms for {word} ({pos}): check the X marks."
            n_inflected_forms_full += len(l_inflections)

            if f"{word}|{pos}" not in d_headword_to_freq:
                d_headword_to_freq[f"{word}|{pos}"] = freq
            else:
                assert d_headword_to_freq[f"{word}|{pos}"] == freq, (word, pos, tup)

            for dic in l_inflections:
                assert dic["word_form"] not in {None, ""}, f"Empty word form for {word} ({pos}): fill in the cell or use 'N/A'."
                word_form = dic["word_form"]
                pos_ud = dic["POS_UD"]
                feats_raw = dic["morphological_features_UD"]

                if feats_raw is not None:
                    l_feats = feats_raw.split("|")
                    assert len(l_feats) == len(set(l_feats)), (word_form, l_feats)  # make sure the list of features does not contain any duplicates
                    feats_sorted = "|".join(sorted(l_feats))
                else:
                    feats_sorted = "N/A"

                l_entries_prov.append((word, pos, word_form, pos_ud, feats_sorted))

    # SORT ITEMS IN FREQUENCY DICTIONARY TO DETERMINE FREQUENCY RANK AND FREQUENCY BAND OF EACH HEADWORD
    l_headwords_sorted_freq = sorted(sorted(
        list(d_headword_to_freq.items()),
        key=lambda x: x[0]),  # second ranking criterion = alphabetically
        key=lambda x: x[1], reverse=True  # first ranking criterion = frequency
    )

    d_headword_to_new_freq = {}

    for idx, tup in enumerate(l_headwords_sorted_freq[:5000]):  # trim resource down to 5,000 entries with highest frequency
        headword = tup[0]
        rank = idx + 1

        if rank <= 1000:
            band = 1
        elif 1000 < rank <= 2000:
            band = 2
        elif 2000 < rank <= 3000:
            band = 3
        elif 3000 < rank <= 4000:
            band = 4
        elif 4000 < rank <= 5000:
            band = 5
        else:
            raise ValueError(f"Invalid rank for {headword}: {rank}.")

        d_headword_to_new_freq[headword] = {"rank": rank, "band": band}

    # CREATE LIST OF FINAL ENTRIES (ADDING ID, FREQUENCY RANK, AND FREQUENCY BAND TO EACH PROVISIONAL ENTRY), SORTED FROM RANK 1 (HIGHEST FREQUENCY) TO RANK 5,000 (LOWEST)
    d_headword_to_id = {headword: 1 for headword in d_headword_to_new_freq}
    l_entries = []
    l_headwords_skipped = []
    n_inflected_forms_skipped = 0

    for tup in l_entries_prov:
        word = tup[0]
        pos = tup[1]

        if f"{word}|{pos}" in d_headword_to_new_freq:
            freq_rank = d_headword_to_new_freq[f"{word}|{pos}"]["rank"]
            freq_band = d_headword_to_new_freq[f"{word}|{pos}"]["band"]
            word_id = f"{freq_rank}-{d_headword_to_id[f'{word}|{pos}']}"
            d_headword_to_id[f"{word}|{pos}"] += 1
            l_entries.append((word_id, freq_rank, freq_band, *tup))
        else:
            l_headwords_skipped.append(f"{word}|{pos}")
            n_inflected_forms_skipped += 1

    assert len(set(l_headwords_skipped)) == 156  # make sure only 156 headwords are skipped (provisional resource contains 5,156 headwords, which needs to be trimmed down to 5,000)
    l_entries_sorted = sorted(l_entries, key=lambda x: x[1])

    # WRITE FINAL RESOURCE TO TSV (FULL AND SAMPLE)
    os.makedirs(config.DIREC_FINAL_RESOURCE, exist_ok=True)
    fn_full = "NLLP-5000.tsv"

    with open(config.DIREC_FINAL_RESOURCE / fn_full, mode="w", newline="", encoding="utf-8") as f_csv:
        csv_writer = csv.writer(f_csv, delimiter="\t", quotechar='"', quoting=csv.QUOTE_MINIMAL)
        csv_writer.writerow([
            "ID", "frequency_rank", "frequency_band", "headword", "POS_Routledge", "inflected_form", "POS_UD",
            "features_UD"
        ])

        for tup in l_entries_sorted:
            csv_writer.writerow(tup)

    fn_sample = "NLLP-5000_sample.tsv"
    random.seed(42)
    set_freq_ranks_random_sample = set(random.sample(range(1, 5001), k=50))  # define which headwords should be included in random sample by randomly choosing 50 frequency ranks (one rank value = one headword)

    with open(config.DIREC_FINAL_RESOURCE / fn_sample, mode="w", newline="", encoding="utf-8") as f_csv:
        csv_writer = csv.writer(f_csv, delimiter="\t", quotechar='"', quoting=csv.QUOTE_MINIMAL)
        csv_writer.writerow([
            "ID", "frequency_rank", "frequency_band", "headword", "POS_Routledge", "inflected_form", "POS_UD",
            "features_UD"
        ])

        for tup in l_entries_sorted:
            freq_rank = tup[1]

            if freq_rank in set_freq_ranks_random_sample:
                csv_writer.writerow(tup)

    # WRITE IMPORTANT MESSAGES TO LOGGER
    logger.info(f"Number of inflected forms (full resource): {n_inflected_forms_full}.")
    logger.info(f"Number of inflected forms (final resource of 5,000 entries): "
                f"{(n_inflected_forms_full - n_inflected_forms_skipped)}.")
